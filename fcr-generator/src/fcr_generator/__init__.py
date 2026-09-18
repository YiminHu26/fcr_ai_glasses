from __future__ import annotations

import argparse
from copy import copy
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO
from math import ceil
from pathlib import Path
from typing import Any
from unicodedata import east_asian_width

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.styles import Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

RAW_SHEET_ALIASES = ("raw data", "raw_data")
TEMPLATE_SHEET_ALIASES = ("fcr_template", "Template")
DEFAULT_TEMPLATE_NAME = "fcr_template.xlsx"
DRAWING_START_ROW = 5
MODULE_START_ROW = 8
MODULES_PER_ROW = 8
TABLE_HEADER_ROW = 14
TABLE_START_ROW = 15
TABLE_END_ROW = 31
CELL_DATE = "G32"
EAST_ASIAN_CHARACTER_WIDTH = 1.3
IMAGE_MAX_ROWS = 3
TABLE_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


@dataclass
class RawRecord:
    bay_id: str
    proj_id: str
    batch_id: str
    hz_drawing: str
    module_symbols: str
    device_id: str
    create_time: str
    inspection_id: str
    inspection_item: str
    inspection_result: str
    inspection_problem: str


def text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def format_report_timestamp() -> str:
    return datetime.now().strftime("%Y-%m-%d_%H-%M-%S")


def merge_inspection_results(records: list[RawRecord]) -> str:
    results = {record.inspection_result.strip().lower() for record in records}
    if "tbd" in results:
        return "TBD"
    if "nok" in results:
        return "NOK"
    if "ok" in results:
        return "OK"
    if "skipped" in results:
        return "Skipped"
    return records[0].inspection_result


def parse_records(ws: Worksheet) -> list[RawRecord]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("raw data sheet is empty")

    columns = {text(value).lower(): index for index, value in enumerate(rows[0])}

    def get(row: tuple[Any, ...], *names: str) -> str:
        for name in names:
            index = columns.get(name.lower())
            if index is not None and index < len(row):
                return text(row[index])
        return ""

    records = []
    for row in rows[1:]:
        if not any(value is not None and text(value) for value in row):
            continue
        records.append(
            RawRecord(
                bay_id=get(row, "bay_id"),
                proj_id=get(row, "proj_id"),
                batch_id=get(row, "batch_id"),
                hz_drawing=get(row, "hz_drawing", "hzdrawing_page", "hz_drawing_page"),
                module_symbols=get(row, "transportation_cell_symbol"),
                device_id=get(row, "device_id"),
                create_time=get(row, "bstudio_create_time", "create_time"),
                inspection_id=get(row, "inspection_id"),
                inspection_item=get(row, "inspection_item"),
                inspection_result=get(row, "inspection_result"),
                inspection_problem=get(row, "inspection_problem"),
            )
        )
    return [record for record in records if record.bay_id]


def copy_template_sheet(source: Worksheet, workbook: Workbook, title: str) -> Worksheet:
    target = workbook.create_sheet(title=title)
    source_dimensions = list(source.column_dimensions.values())
    grouped_dimension = next(
        (
            item
            for item in source_dimensions
            if item.min is not None and item.max is not None and item.min == 1
        ),
        None,
    )
    for column in range(1, source.max_column + 1):
        column_letter = get_column_letter(column)
        dimension = next(
            (
                item
                for item in source_dimensions
                if item.min is not None
                and item.max is not None
                and item.min <= column <= item.max
            ),
            source.column_dimensions[column_letter]
            if column_letter in source.column_dimensions
            else grouped_dimension,
        )
        if dimension is not None and dimension.width is not None:
            target.column_dimensions[column_letter].width = dimension.width
        if dimension is not None:
            target.column_dimensions[column_letter].hidden = dimension.hidden
            target.column_dimensions[column_letter].outlineLevel = dimension.outlineLevel
            target.column_dimensions[column_letter].collapsed = dimension.collapsed
    for key, dimension in source.row_dimensions.items():
        target.row_dimensions[key].height = dimension.height
    for row in source.iter_rows():
        for source_cell in row:
            target_cell = target.cell(source_cell.row, source_cell.column)
            target_cell.value = source_cell.value
            if source_cell.has_style:
                target_cell.font = copy(source_cell.font)
                target_cell.border = copy(source_cell.border)
                target_cell.fill = copy(source_cell.fill)
                target_cell.alignment = copy(source_cell.alignment)
                target_cell.protection = copy(source_cell.protection)
            if source_cell.number_format:
                target_cell.number_format = source_cell.number_format
    for merged_range in source.merged_cells.ranges:
        target.merge_cells(str(merged_range))
    for source_image in source._images:
        image = OpenpyxlImage(BytesIO(source_image._data()))
        max_height_points = sum(
            source.row_dimensions[row].height or 15
            for row in range(1, IMAGE_MAX_ROWS + 1)
        )
        max_height_pixels = max_height_points * 96 / 72
        scale = min(1, max_height_pixels / source_image.height)
        image.height = round(source_image.height * scale)
        image.width = round(source_image.width * scale)
        anchor = source_image.anchor
        if isinstance(anchor, str):
            image.anchor = anchor
        else:
            image.anchor = f"{get_column_letter(anchor._from.col + 1)}{anchor._from.row + 1}"
        target.add_image(image)
    return target


def write_value(ws: Worksheet, coordinate: str, value: Any, reference: Any) -> None:
    cell = ws[coordinate]
    cell.value = value
    if reference is not None and reference.has_style:
        cell.font = copy(reference.font)
        cell.border = copy(reference.border)
        cell.fill = copy(reference.fill)
        cell.alignment = copy(reference.alignment)
        cell.protection = copy(reference.protection)


def adjust_inspection_row_height(ws: Worksheet, row: int) -> None:
    """根据检测项和备注文本长度调整合并单元格所在行的高度。"""
    def merged_width(start_column: int, end_column: int) -> float:
        return sum(
            ws.column_dimensions[get_column_letter(column)].width or 9
            for column in range(start_column, end_column + 1)
        )

    def line_count(value: Any, width: float) -> int:
        value = text(value)
        if not value:
            return 1

        def display_width(line: str) -> int:
            return sum(
                EAST_ASIAN_CHARACTER_WIDTH
                if east_asian_width(character) in ("W", "F")
                else 1
                for character in line
            )

        characters_per_line = max(1, width)
        return sum(
            max(1, ceil(display_width(line) / characters_per_line))
            for line in value.splitlines()
        )

    for coordinate, start_column, end_column in (
        (f"B{row}", 2, 5),
        (f"G{row}", 7, 8),
    ):
        cell = ws[coordinate]
        alignment = copy(cell.alignment)
        alignment.wrap_text = True
        alignment.vertical = "center"
        cell.alignment = alignment
        cell_lines = line_count(cell.value, merged_width(start_column, end_column))
        row_height = {
            1: 25,
            2: 30,
            3: 40,
        }.get(cell_lines, 40 + (cell_lines - 3) * 10)
        ws.row_dimensions[row].height = max(ws.row_dimensions[row].height or 0, row_height)


def apply_table_borders(ws: Worksheet) -> None:
    """为表头到 A 列最后一个有内容行的 A:H 单元格添加边框。"""
    last_row = max(
        (
            row
            for row in range(TABLE_HEADER_ROW, TABLE_END_ROW + 1)
            if ws.cell(row=row, column=1).value not in (None, "")
        ),
        default=TABLE_HEADER_ROW,
    )
    for row in range(TABLE_HEADER_ROW, last_row + 1):
        for column in range(1, 9):
            ws.cell(row=row, column=column).border = copy(TABLE_BORDER)


def fill_header(ws: Worksheet, template: Worksheet, records: list[RawRecord]) -> None:
    first = records[0]
    write_value(ws, "B4", first.device_id, template["A4"])
    write_value(ws, "D4", first.proj_id, template["C4"])
    write_value(ws, "H4", first.batch_id, template["G4"])

    drawings = []
    seen_drawings = set()
    modules = []
    seen_modules = set()
    for record in records:
        if record.hz_drawing and record.hz_drawing not in seen_drawings:
            seen_drawings.add(record.hz_drawing)
            drawings.append(record.hz_drawing)
        for module in record.module_symbols.split("+"):
            module = module.strip()
            if module and module not in seen_modules:
                seen_modules.add(module)
                modules.append(module)

    for index, drawing in enumerate(drawings):
        write_value(ws, f"B{DRAWING_START_ROW + index}", drawing, template["A6"])

    for index, module in enumerate(modules):
        row = MODULE_START_ROW + index // MODULES_PER_ROW
        column = 1 + index % MODULES_PER_ROW
        write_value(ws, f"{get_column_letter(column)}{row}", module, template["A9"])

    last_record = next(
        (record for record in reversed(records) if record.create_time),
        None,
    )
    if last_record:
        write_value(ws, CELL_DATE, last_record.create_time[:10], template["F32"])


def fill_inspection_table(ws: Worksheet, template: Worksheet, records: list[RawRecord]) -> None:
    grouped_records: dict[str, list[RawRecord]] = {}
    for record in records:
        grouped_records.setdefault(record.inspection_item, []).append(record)
    merged_records = list(grouped_records.values())

    if len(merged_records) > TABLE_END_ROW - TABLE_START_ROW + 1:
        raise ValueError(
            f"bay {records[0].bay_id} has {len(merged_records)} inspection rows after merging, "
            f"but the template supports at most {TABLE_END_ROW - TABLE_START_ROW + 1} rows"
        )

    for row in range(TABLE_START_ROW, TABLE_END_ROW + 1):
        for column in (1, 2, 6, 7):
            ws.cell(row=row, column=column).value = None

    for offset, same_item_records in enumerate(merged_records):
        row = TABLE_START_ROW + offset
        first = same_item_records[0]
        problem_parts = []
        for record in same_item_records:
            if record.inspection_result.strip().upper() != "OK" and (
                record.module_symbols or record.inspection_problem
            ):
                problem_parts.append(f"{record.module_symbols}{record.inspection_problem}")
        values = {
            1: first.inspection_id,
            2: first.inspection_item,
            6: merge_inspection_results(same_item_records),
            7: "；".join(problem_parts),
        }
        for column, value in values.items():
            coordinate = f"{get_column_letter(column)}{row}"
            write_value(ws, coordinate, value, template.cell(TABLE_START_ROW, column))
        adjust_inspection_row_height(ws, row)

    adjust_inspection_row_height(ws, TABLE_HEADER_ROW)
    apply_table_borders(ws)


def generate_report(
    input_path: Path,
    output_path: Path | None = None,
    template_path: Path | None = None,
) -> Path:
    data_workbook = load_workbook(input_path, data_only=True)
    raw_sheet_name = next(
        (name for name in RAW_SHEET_ALIASES if name in data_workbook.sheetnames),
        data_workbook.sheetnames[0] if len(data_workbook.sheetnames) == 1 else None,
    )
    if raw_sheet_name is None:
        raise ValueError(
            f"missing raw data sheet; expected one of: {RAW_SHEET_ALIASES}"
        )

    records = parse_records(data_workbook[raw_sheet_name])
    data_workbook.close()
    records_by_bay: dict[str, list[RawRecord]] = {}
    for record in records:
        records_by_bay.setdefault(record.bay_id, []).append(record)
    if not records_by_bay:
        raise ValueError("no records with bay_id were found")

    if output_path is None:
        first_record = records[0]
        report_date = format_report_timestamp()
        output_path = input_path.parent / (
            f"FCR_{first_record.proj_id}_{first_record.bay_id}_{report_date}.xlsx"
        )

    template_path = template_path or input_path.parent / DEFAULT_TEMPLATE_NAME
    if not template_path.exists():
        raise FileNotFoundError(f"template file not found: {template_path}")
    template_workbook = load_workbook(template_path, data_only=False)
    template_sheet_name = next(
        (
            name
            for name in TEMPLATE_SHEET_ALIASES
            if name in template_workbook.sheetnames
        ),
        template_workbook.sheetnames[0]
        if len(template_workbook.sheetnames) == 1
        else None,
    )
    if template_sheet_name is None:
        raise ValueError(
            f"missing template sheet; expected one of: {TEMPLATE_SHEET_ALIASES}"
        )

    template = template_workbook[template_sheet_name]
    output = Workbook()
    output.remove(output.active)
    for bay_id, bay_records in records_by_bay.items():
        sheet = copy_template_sheet(template, output, bay_id)
        fill_header(sheet, template, bay_records)
        fill_inspection_table(sheet, template, bay_records)

    output.save(output_path)
    template_workbook.close()
    print(f"Generated {len(records_by_bay)} report sheet(s): {output_path}")
    return output_path


def main() -> None:
    parser = argparse.ArgumentParser(description="Fill FCR report templates from raw Excel data")
    parser.add_argument("input", type=Path, help="source Excel workbook")
    parser.add_argument("-o", "--output", type=Path)
    parser.add_argument("--template", type=Path, help="template workbook; defaults to fcr_template.xlsx")
    args = parser.parse_args()
    generate_report(args.input, args.output, args.template)


if __name__ == "__main__":
    main()